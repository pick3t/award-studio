"""生成测试数据：随机成绩表 + 奖状模板。

用法：
    python tools/make_samples.py [输出目录]
默认输出到项目根目录的 samples/。
"""

from __future__ import annotations

import random
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SURNAMES = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜"
GIVEN = ["子轩", "雨桐", "睿泽", "梓涵", "思远", "欣怡", "浩然", "嘉怡", "宇轩", "若曦",
         "明轩", "佳琪", "天佑", "诗涵", "俊杰", "梦瑶", "志远", "雅静", "博文", "晨曦",
         "逸凡", "可欣", "承宇", "静怡", "文昊", "书瑶", "皓然", "语彤", "锦程", "清越"]
CLASSES = ["高一(1)班", "高一(3)班", "高二(2)班", "高二(5)班", "高三(1)班", "高三(4)班",
           "计算机 2301", "计算机 2302", "电子 2301", "机械 2302"]
PROJECTS = ["智能垃圾分类装置", "校园能耗监测系统", "基于视觉的无人小车", "低成本水质检测仪",
            "助老跌倒报警器", "便携式太阳能充电站", "语音控制机械臂", "AI 作文批改助手",
            "智能温室环境调控", "共享单车停放引导系统", "手势识别钢琴", "宿舍用电安全监测"]


# ------------------------------------------------------------------ Excel
def build_scores(count: int = 32, seed: int = 2026) -> list[dict]:
    rng = random.Random(seed)
    used: set[str] = set()
    rows = []
    for index in range(1, count + 1):
        while True:
            name = rng.choice(SURNAMES) + rng.choice(GIVEN)
            if name not in used:
                used.add(name)
                break
        level = min(99.0, max(55.0, rng.gauss(84, 8)))
        judges = [round(min(100.0, max(40.0, level + rng.gauss(0, 3.2))), 1) for _ in range(3)]
        rows.append(
            {
                "no": f"C{index:03d}",
                "name": name,
                "class": rng.choice(CLASSES),
                "project": rng.choice(PROJECTS),
                "judges": judges,
            }
        )
    return rows


def make_excel(path: Path, count: int = 32, seed: int = 2026) -> Path:
    rows = build_scores(count, seed)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "决赛评分"

    headers = ["编号", "姓名", "班级", "作品名称", "评委一", "评委二", "评委三", "总分"]
    widths = [10, 11, 14, 30, 10, 10, 10, 10]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    # 标题区
    sheet.merge_cells("A1:H1")
    title = sheet["A1"]
    title.value = "2026 校园科技创新大赛 · 决赛评分表"
    title.font = Font(name="微软雅黑", size=18, bold=True, color="1F3864")
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells("A2:H2")
    subtitle = sheet["A2"]
    subtitle.value = "评分日期：2026-10-20   评分方式：百分制（三位评委打分，总分自动汇总）   制表：大赛组委会"
    subtitle.font = Font(name="微软雅黑", size=9, color="6B7280")
    subtitle.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 20

    header_row = 4
    header_fill = PatternFill("solid", fgColor="1F3864")
    thin = Side(style="thin", color="D6DCE5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, name in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=name)
        cell.font = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    sheet.row_dimensions[header_row].height = 26

    zebra = PatternFill("solid", fgColor="F4F7FB")
    for offset, row in enumerate(rows):
        excel_row = header_row + 1 + offset
        values = [row["no"], row["name"], row["class"], row["project"],
                  row["judges"][0], row["judges"][1], row["judges"][2]]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=column, value=value)
            cell.font = Font(name="微软雅黑", size=10.5, color="1F2937")
            cell.alignment = Alignment(
                horizontal="left" if column == 4 else "center", vertical="center"
            )
            cell.border = border
            if offset % 2 == 1:
                cell.fill = zebra
        total_cell = sheet.cell(row=excel_row, column=8, value=f"=SUM(E{excel_row}:G{excel_row})")
        total_cell.font = Font(name="微软雅黑", size=10.5, bold=True, color="1F3864")
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_cell.border = border
        total_cell.number_format = "0.0"
        if offset % 2 == 1:
            total_cell.fill = zebra
        sheet.row_dimensions[excel_row].height = 22

    last_row = header_row + len(rows)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:H{last_row}"

    # 说明页：告诉使用者这张表可以直接喂给软件
    guide = workbook.create_sheet("填写说明")
    guide.column_dimensions["A"].width = 96
    notes = [
        ("这是一个测试用成绩表，共 %d 名选手。" % len(rows), True),
        ("", False),
        ("· 软件会自动识别表头，不要求表头在第 1 行；标题行、说明行会被跳过。", False),
        ("· 姓名列可写作 姓名 / 选手 / 参赛者；分数列可写作 总分 / 成绩 / 得分 / 评分。", False),
        ("· 这里的「总分」是 SUM 公式，软件两种都能读：Excel 保存过的缓存值，或公式自动求和。", False),
        ("· 若没有「总分」列，但有多个评委列，软件会把它们相加。", False),
        ("· 「班级」「作品名称」等列会自动成为奖状里可用的字段，例如 {{作品名称}}。", False),
        ("", False),
        ("换成真实比赛数据时，保留同样的列结构即可，行数不限。", False),
    ]
    for index, (text, bold) in enumerate(notes, start=1):
        cell = guide.cell(row=index, column=1, value=text)
        cell.font = Font(name="微软雅黑", size=12 if bold else 10.5, bold=bold, color="1F2937")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        guide.row_dimensions[index].height = 26 if bold else 22

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


# ------------------------------------------------------------------ Word 模板
def _set_font(run, name: str, size: float, bold: bool = False,
              color: str | None = None, eastasia: str | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), eastasia or name)


def _char_spacing(run, twips: int) -> None:
    rpr = run._element.get_or_add_rPr()
    element = OxmlElement("w:spacing")
    element.set(qn("w:val"), str(twips))
    rpr.insert_element_before(
        element, "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight",
        "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
        "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
    )


def _page_borders(section, color: str = "B08D57") -> None:
    sect_pr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "double")
        element.set(qn("w:sz"), "18")
        element.set(qn("w:space"), "24")
        element.set(qn("w:color"), color)
        borders.append(element)
    sect_pr.insert_element_before(
        borders, "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt", "w:vAlign",
        "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter",
        "w:docGrid", "w:printerSettings", "w:sectPrChange",
    )


def _paragraph_border_bottom(paragraph, color: str = "B08D57", size: int = 6) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    ppr.insert_element_before(
        p_bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
        "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
        "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def make_certificate_template(path: Path) -> Path:
    document = Document()

    # 页面：A4 横向
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.2)
    section.right_margin = Cm(3.2)
    _page_borders(section)

    # 默认字体
    normal = document.styles["Normal"]
    normal.font.name = "楷体"
    normal.font.size = Pt(16)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")

    # 证书编号（右上角）
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("证书编号：AS-{{编号}}")
    _set_font(run, "楷体", 10.5, color="8C8C8C")
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(26)

    # 主标题
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("获 奖 证 书")
    _set_font(run, "Times New Roman", 42, bold=True, color="9C2B2B", eastasia="宋体")
    _char_spacing(run, 160)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.2

    # 英文副标题
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CERTIFICATE  OF  AWARD")
    _set_font(run, "Times New Roman", 11, color="B08D57")
    _char_spacing(run, 60)
    p.paragraph_format.space_after = Pt(40)

    # 正文
    body_format = {
        "line_spacing": 2.4,
        "first_line_indent": Pt(32),
        "space_after": Pt(10),
    }

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run("兹有　")
    _set_font(run, "楷体", 18, color="1A1A1A")
    run = p.add_run("{{姓名}}")
    _set_font(run, "黑体", 20, bold=True, color="1A1A1A", eastasia="黑体")
    run.font.underline = True
    run = p.add_run("　同学，在「{{比赛名称}}」中表现优异，")
    _set_font(run, "楷体", 18, color="1A1A1A")
    for key, value in body_format.items():
        setattr(p.paragraph_format, key, value)
    p.paragraph_format.space_after = Pt(0)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run("荣获 ")
    _set_font(run, "楷体", 18, color="1A1A1A")
    run = p.add_run("{{奖项}}")
    _set_font(run, "黑体", 20, bold=True, color="9C2B2B", eastasia="黑体")
    run = p.add_run("。特发此证，以资鼓励。")
    _set_font(run, "楷体", 18, color="1A1A1A")
    for key, value in body_format.items():
        setattr(p.paragraph_format, key, value)

    # 装饰分隔
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("◆　◆　◆")
    _set_font(run, "宋体", 11, color="B08D57")
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(64)

    # 落款
    for text in ("{{比赛名称}} 组织委员会", "{{日期}}"):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        _set_font(run, "楷体", 15, color="1A1A1A")
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.right_indent = Cm(1.2)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def main() -> None:
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "samples"
    excel = make_excel(out_dir / "决赛评分表.xlsx")
    template = make_certificate_template(out_dir / "奖状模板.docx")
    print(f"成绩表：{excel}  ({excel.stat().st_size / 1024:.1f} KB)")
    print(f"奖状模板：{template}  ({template.stat().st_size / 1024:.1f} KB)")
    print(f"生成日期：{date.today().isoformat()}")


if __name__ == "__main__":
    main()
