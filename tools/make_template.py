"""生成竖版奖状模板（含项目组、参赛项目占位符）。

    python tools/make_template.py [输出路径]

默认写到 samples/奖状模板_项目组.docx，并复制一份到桌面方便试用。
占位符：{{比赛名称}} {{项目组}} {{参赛项目}} {{代表队}} {{姓名}} {{奖项}}
        {{名次}} {{颁发单位}} {{日期}} {{编号}}
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))

from make_samples import _set_font  # noqa: E402

INK = "1A1A1A"
RED = "9C2B2B"
GOLD = "A07C3C"
GREY = "6B7280"


def setup_page(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "12")
        element.set(qn("w:space"), "24")
        element.set(qn("w:color"), GOLD)
        borders.append(element)
    section._sectPr.append(borders)


def para(document: Document, text: str, size: float, *, align=WD_ALIGN_PARAGRAPH.CENTER,
         bold=False, color=INK, before: float = 0, after: float = 0,
         eastasia: str = "宋体", spaced: float = 0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    latin = "Times New Roman" if eastasia in ("宋体", "楷体", "黑体") else eastasia
    _set_font(run, latin, size, bold=bold, color=color, eastasia=eastasia)
    if spaced:
        for ch in run.text:
            pass
        run.font.spacing = Pt(spaced)


def hr(document: Document, width_pct: int = 60, color: str = GOLD) -> None:
    """一条居中的细线，用下边框实现。"""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)
    ind = OxmlElement("w:ind")
    margin = int((100 - width_pct) / 2 * 56.7)
    ind.set(qn("w:left"), str(margin))
    ind.set(qn("w:right"), str(margin))
    p_pr.append(ind)


def build(path: Path) -> Path:
    document = Document()
    setup_page(document)

    para(document, "", 11, after=40)
    para(document, "{{比赛名称}}", 11, color=GREY, eastasia="楷体", after=24)
    para(document, "获 奖 证 书", 34, bold=True, color=RED, after=2)
    para(document, "CERTIFICATE  OF  AWARD", 9.5, color=GOLD, eastasia="Times New Roman", after=6)
    hr(document, 62)
    para(document, "", 8, after=4)

    para(document, "项目组：{{项目组}}", 11, color=GREY, eastasia="楷体", after=2)
    para(document, "参赛项目：{{参赛项目}}　　参赛单位：{{代表队}}", 11, color=GREY,
         eastasia="楷体", after=34)

    para(document, "兹 有", 15, eastasia="楷体", after=8)
    para(document, "{{姓名}}", 26, bold=True, color=RED, after=8)
    para(document, "同 学", 15, eastasia="楷体", after=22)
    para(document, "在本次比赛中表现优异，荣获", 15, eastasia="楷体", after=8)
    para(document, "{{奖项}}", 22, bold=True, color=RED, after=4)
    para(document, "（第 {{名次}} 名）", 11, color=GREY, eastasia="楷体", after=30)
    para(document, "特发此证，以资鼓励", 15, eastasia="楷体", after=96)

    para(document, "{{颁发单位}}", 13, align=WD_ALIGN_PARAGRAPH.RIGHT, eastasia="楷体", after=4)
    para(document, "{{日期}}", 13, align=WD_ALIGN_PARAGRAPH.RIGHT, eastasia="楷体", after=54)
    para(document, "证书编号：{{编号}}", 10, align=WD_ALIGN_PARAGRAPH.LEFT, color=GREY)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def main() -> None:
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "samples" / "奖状模板_项目组.docx"
    build(target)
    print(f"saved {target}")
    desktop = Path.home() / "Desktop" / target.name
    try:
        shutil.copy(target, desktop)
        print(f"saved {desktop}")
    except OSError as exc:
        print(f"复制到桌面失败（{exc}），不影响使用")


if __name__ == "__main__":
    main()
