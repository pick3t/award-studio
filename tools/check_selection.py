"""自测：缺分数也能加载、缺分的人排末尾、只对选中的人排名。

    python tools/check_selection.py

验证点：
1. 分数列整列为空也能加载，不弹「读取失败」
2. 勾选分数列后：有分的正常排名，没分的排到最后、名次「—」、不获奖
3. 提示里写明有几个人缺分数
4. 只选 3 个人参与排名时，其余人排到末尾不参与
5. 取消筛选后恢复全员排名
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import app as gui  # noqa: E402

SAMPLE = ROOT / "samples" / "真实场地" / "一号场地成绩单_模拟分数.xlsx"
BLANK = Path.home() / "Desktop" / "一号场地成绩单.xlsx"
SHEET = "一场地4组"

failures: list[str] = []


def check(ok: bool, message: str) -> None:
    print(("  [OK] " if ok else "  [!!] ") + message)
    if not ok:
        failures.append(message)


def cells(tree) -> list[list[str]]:
    return [list(tree.item(item, "values")) for item in tree.get_children()]


def make_partial() -> Path:
    """复制一份样本，把前 3 行分数抹掉。"""
    target = Path(r"C:\Users\admin\AppData\Local\Temp\_partial_scores.xlsx")
    workbook = load_workbook(SAMPLE)
    sheet = workbook[SHEET]
    for row in (5, 6, 7):
        sheet.cell(row=row, column=6).value = None
    workbook.save(target)
    return target


def main() -> int:
    window = gui.AwardStudio()
    window.withdraw()
    try:
        # 1) 分数整列为空也要能加载
        if BLANK.exists():
            window._load_excel(BLANK)
            window.update_idletasks()
            check(bool(window.players), f"空分数表能加载，读到 {len(window.players)} 人")
            window._on_heading_click("分数")
            ranks = [entry.rank for entry in window.entries]
            check(all(rank == 0 for rank in ranks), "全部人没有分数时都不排名")
            check("缺分数" in window.table_note.cget("text"),
                  f"提示里写明缺分数：{window.table_note.cget('text')}")
        else:
            print(f"  [跳过] 找不到 {BLANK}")

        # 2) 部分人缺分数
        partial = make_partial()
        window._load_excel(partial)
        window._apply_sheet(SHEET, keep_columns=True)
        window._on_heading_click("分数")
        window.update_idletasks()

        entries = window.entries
        missing = [e for e in entries if not e.player.scored]
        scored = [e for e in entries if e.player.scored]
        check(len(missing) == 3, f"识别出 3 人缺分数（实际 {len(missing)}）")
        check(all(e.rank == 0 for e in entries[len(scored):]), "缺分数的人排在最后且没有名次")
        check(entries[0].rank == 1 and entries[len(scored) - 1].rank != 0, "有分数的人正常排名")
        check(all(not e.is_winner for e in missing), "缺分数的人不参与发奖")
        check("3 人缺分数" in window.table_note.cget("text"),
              f"提示写明缺分数人数：{window.table_note.cget('text')}")
        row_values = cells(window.tree)
        check(row_values[-1][0] == "☑", "末行仍是勾选状态（默认全员参与）")

        # 3) 只对选中的人排名
        total = len(window.entries)
        window._select_range(0, 2)
        window._after_pick_change()
        ranks = [entry.rank for entry in window.entries]
        check(len(window.picked) == 3, f"拖选后选中 3 人（{len(window.picked)}）")
        check(ranks[:3] == [1, 2, 3], f"选中的人重新排名：{ranks[:3]}")
        check(all(rank == 0 for rank in ranks[3:]), "未选中的人没有名次")
        check(len(window.entries) == total, "总行数不变，未选中的人排在末尾")
        check("仅 3 人参与排名" in window.table_note.cget("text"),
              f"提示写明参与人数：{window.table_note.cget('text')}")
        awards = [e for e in window.entries if e.is_winner]
        check(len(awards) <= 3, f"获奖人数不超过选中的 3 人（{len(awards)}）")

        # 4) 表头「全选 / 取消」来回切换
        window._toggle_all()
        ranks = [entry.rank for entry in window.entries]
        check(not window.filter_on, "点表头「全选」后恢复全员参与")
        check(ranks[0] == 1 and any(rank == 0 for rank in ranks), "恢复全员排名（缺分数的仍排末尾）")
        check(window.tree.heading("pick")["text"] == "取消", "全员时表头显示「取消」")

        window._toggle_all()
        window.update_idletasks()
        check(window.filter_on and not window.picked, "再点表头「取消」后一个人都不参与")
        check(window.tree.heading("pick")["text"] == "全选", "无人参与时表头显示「全选」")
        check(all(entry.rank == 0 for entry in window.entries), "没人参与时全部没有名次")
        check(str(window.generate_btn.cget("state")) == "disabled", "没人参与时生成按钮不可用")

        window._toggle_all()
        window.update_idletasks()
        check(window._all_picked() and window.entries[0].rank == 1, "再点「全选」恢复排名")
    finally:
        window.destroy()

    print()
    if failures:
        print(f"失败 {len(failures)} 项：")
        for item in failures:
            print(" - " + item)
        return 1
    print("全部通过 ✓ 缺分数可加载、缺分者排末尾、可按勾选部分排名")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
