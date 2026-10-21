"""端到端自测：读成绩表 -> 排名 -> 生成奖状 -> 校验内容。

不需要图形界面，可在 CI 或命令行直接运行：
    python tools/smoke_test.py
"""

from __future__ import annotations

import shutil
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

from core.certificate import (  # noqa: E402
    CertificateContext,
    generate_all,
    inspect_template,
)
from core.excel_reader import load_players  # noqa: E402
from core.ranking import rank_players, winners_of  # noqa: E402

SAMPLES = ROOT / "samples"
OUTPUT = ROOT / "output" / "自测输出"


def document_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def main() -> int:
    failures: list[str] = []

    def check(condition: bool, message: str) -> None:
        print(("  [OK] " if condition else "  [!!] ") + message)
        if not condition:
            failures.append(message)

    excel = SAMPLES / "决赛评分表.xlsx"
    template = SAMPLES / "奖状模板.docx"
    for path in (excel, template):
        if not path.exists():
            print(f"缺少测试文件：{path}，请先运行 tools/make_samples.py")
            return 2

    print("1) 解析成绩表")
    players, warnings, report = load_players(excel)
    print(f"   工作表={report.sheet} 表头行={report.header_row} -> {report.describe()}")
    for warning in warnings:
        print("   ! " + warning)
    check(len(players) == 32, f"读到 32 名选手（实际 {len(players)}）")
    check(all(p.score > 0 for p in players), "所有分数都成功解析（含 SUM 公式兜底）")
    check(60 <= max(p.score for p in players) <= 300, "分数在合理区间")

    print("2) 排名")
    entries = rank_players(players, winners=5)
    scores = [e.player.score for e in entries]
    check(scores == sorted(scores, reverse=True), "按分数降序排列")
    check(entries[0].rank == 1, "第一名名次为 1")
    ties = [(e.rank, e.player.score) for e in entries]
    consistent = all(
        ties[i][0] == ties[i - 1][0]
        for i in range(1, len(ties))
        if abs(ties[i][1] - ties[i - 1][1]) < 1e-9
    )
    check(consistent, "同分并列名次一致")

    winners = winners_of(entries, 5)
    check(len(winners) == 5, f"取到 5 名获奖者（实际 {len(winners)}）")
    print("   前 5 名：" + "、".join(f"{e.rank}.{e.player.name}({e.player.score:g})" for e in winners))

    print("3) 奖状模板")
    keys = inspect_template(template)
    check({"姓名", "名次"} <= keys, f"模板包含 姓名/名次 占位符：{sorted(keys)}")

    print("4) 生成奖状")
    if OUTPUT.exists():
        shutil.rmtree(OUTPUT)
    context = CertificateContext(
        contest="2026 校园科技创新大赛",
        issuer="大赛组委会",
        date=date.today().strftime("%Y 年 %m 月 %d 日"),
    )
    created = generate_all(
        template=template,
        output_dir=OUTPUT,
        winners=winners,
        context=context,
        merge=True,
    )
    single = [p for p in created if "合并" not in p.name]
    merged = [p for p in created if "合并" in p.name]
    check(len(single) == 5, f"生成 5 份单人奖状（实际 {len(single)}）")
    check(len(merged) == 1, "生成 1 份合并版")

    for path, entry in zip(single, winners):
        text = document_text(path)
        ok_name = entry.player.name in text
        ok_rank = f"第 {entry.rank} 名" in text
        ok_empty = "{{" not in text and "}}" not in text
        check(ok_name and ok_rank and ok_empty,
              f"{path.name} 姓名/名次已填入且无残留占位符")
        print(f"       {text.strip().splitlines()[:3]}")

    if merged:
        text = document_text(merged[0])
        pages = text.count("\f") + len([p for p in Document(str(merged[0])).paragraphs if "w:br" in p._p.xml])
        check(all(e.player.name in text for e in winners), "合并版包含全部获奖者姓名")
        check(text.count("获 奖 证 书") == 5, f"合并版共 5 页（证书标题出现 {text.count('获 奖 证 书')} 次）")

    print()
    if failures:
        print(f"失败 {len(failures)} 项：")
        for item in failures:
            print(" - " + item)
        return 1
    print("全部通过 ✓  输出目录：" + str(OUTPUT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
