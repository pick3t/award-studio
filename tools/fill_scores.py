"""给成绩表填模拟分数，并顺手规整团体赛的队员名单。

两件事：
1. 每个工作表的得分列填 8.30–9.90 的随机两位小数
2. 「参赛选手」列规整成「单个空格分隔、两字姓名中间不带空格」：
   * `张  锐`（两字名被空格劈开）→ `张锐`
   * `徐泽宁王琦超`（两人粘在一起）→ `徐泽宁 王琦超`

结果另存为 <原名>_模拟分数.xlsx，原文件不动。

    python tools/fill_scores.py                # 处理桌面上的三张场地成绩单
    python tools/fill_scores.py 文件1.xlsx ...  # 处理指定文件
"""

from __future__ import annotations

import random
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

MEMBER_HEADERS = ("参赛选手", "队员", "成员", "选手名单")

DESKTOP = Path.home() / "Desktop"
DEFAULT_FILES = ["一号场地成绩单.xlsx", "二号场地成绩单.xlsx", "三号场地成绩单.xlsx"]
SCORE_HINTS = ("得分", "成绩", "分数", "总分", "评")
NAME_HINTS = ("姓名", "选手", "参赛者", "队员", "代表队", "队名")
SUMMARY_HINTS = ("名次", "备注", "序号", "队", "项目")


def _norm(value) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\n", "").replace("\r", "")
    return re.sub(r"\s+", "", text)


FOOTER_HINTS = ("仲裁", "裁判", "签名", "签字", "记录", "统计", "合计", "制表", "教练")


def find_header(ws) -> tuple[int, int, int, int] | None:
    """返回 (表头行, 得分列, 数据起始行, 名称列)；两行表头时数据从表头行 + 2 开始。"""
    for row in range(1, min(ws.max_row, 12) + 1):
        texts = [_norm(ws.cell(row=row, column=c).value) for c in range(1, ws.max_column + 1)]
        has_name = any(any(h in t for h in NAME_HINTS) for t in texts)
        if not has_name:
            continue
        name_col = next(i for i, t in enumerate(texts)
                        if any(h in t for h in NAME_HINTS)) + 1
        for idx, text in enumerate(texts):
            if not text or any(h in text for h in SUMMARY_HINTS):
                continue
            if any(h in text for h in SCORE_HINTS):
                return row, idx + 1, row + 1, name_col
        # 「最后」在上一行、「得分」在下一行：合并后再找一次
        if row + 1 <= ws.max_row:
            nxt = [_norm(ws.cell(row=row + 1, column=c).value)
                   for c in range(1, ws.max_column + 1)]
            for idx, text in enumerate(texts):
                merged = text + (nxt[idx] if idx < len(nxt) else "")
                if not merged or any(h in merged for h in SUMMARY_HINTS):
                    continue
                if any(h in merged for h in SCORE_HINTS):
                    return row, idx + 1, row + 2, name_col
    return None


def fix_roster(text: str) -> tuple[str, list[str]]:
    """规整队员名单，返回 (新文本, 改动说明)。

    规则只有两条，针对已知的两种录入问题：
    * 相邻的两个单字 → 一个两字姓名（张  锐 → 张锐）
    * 六个汉字连写 → 两个三字姓名（徐泽宁王琦超 → 徐泽宁 王琦超）
    """
    tokens = str(text or "").replace("\n", " ").split()
    fixed: list[str] = []
    notes: list[str] = []
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if len(token) == 1 and index + 1 < len(tokens) and len(tokens[index + 1]) == 1:
            merged = token + tokens[index + 1]
            notes.append(f"{token} {tokens[index + 1]} → {merged}")
            fixed.append(merged)
            index += 2
            continue
        if len(token) == 6:
            left, right = token[:3], token[3:]
            notes.append(f"{token} → {left} {right}")
            fixed.extend([left, right])
            index += 1
            continue
        fixed.append(token)
        index += 1
    return " ".join(fixed), notes


def fill(path: Path, out_dir: Path | None = None) -> Path:
    workbook = load_workbook(path)
    rng = random.Random(f"{path.name}-2026")
    report = []
    for ws in workbook.worksheets:
        found = find_header(ws)
        if not found:
            report.append(f"  [{ws.title}] 没找到得分列，跳过")
            continue
        header_row, col, start, name_col = found
        header_texts = [_norm(ws.cell(row=header_row, column=c).value)
                        for c in range(1, ws.max_column + 1)]
        member_col = next((i + 1 for i, t in enumerate(header_texts)
                           if any(h in t for h in MEMBER_HEADERS)), None)
        roster_notes: list[str] = []
        roster_rows = 0
        if member_col:
            for row in range(start, ws.max_row + 1):
                cell = ws.cell(row=row, column=member_col)
                if not isinstance(cell.value, str):
                    continue
                new_text, notes = fix_roster(cell.value)
                if notes:
                    cell.value = new_text
                    roster_rows += 1
                    roster_notes.extend(notes)
        filled = 0
        for row in range(start, ws.max_row + 1):
            name = _norm(ws.cell(row=row, column=name_col).value)
            # 没有姓名 / 代表队的是空行或表尾签名行，不填分
            if not name or any(h in name for h in FOOTER_HINTS):
                continue
            ws.cell(row=row, column=col).value = round(rng.uniform(8.30, 9.90), 2)
            filled += 1
        line = f"  [{ws.title}] 得分列第 {col} 列（表头第 {header_row} 行），填入 {filled} 行"
        if roster_rows:
            line += f"；规整队员名单 {roster_rows} 行"
        report.append(line)
        if roster_notes:
            for note in roster_notes:
                report.append(f"      · {note}")

    target_dir = out_dir or path.parent
    target = target_dir / f"{path.stem}_模拟分数.xlsx"
    workbook.save(target)
    print(f"{path.name} → {target.name}")
    print("\n".join(report))
    return target


def main() -> None:
    args = sys.argv[1:]
    files = [Path(a) for a in args] if args else [DESKTOP / n for n in DEFAULT_FILES]
    out_dir = Path(r"D:\WorkBuddyCNWorkspace\Code\AwardStudio\samples\真实场地") if not args else None
    if out_dir:
        out_dir.mkdir(parents=True, exist_ok=True)
    for path in files:
        if not path.exists():
            print(f"找不到 {path}")
            continue
        fill(path, out_dir)
        print()


if __name__ == "__main__":
    main()
